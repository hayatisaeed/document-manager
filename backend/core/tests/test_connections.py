"""Drawings, spreadsheets, links between files, the link graph, comments and the commit tree."""
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core.services import sheets

from .test_api import BaseTest

INTRO = "manuscript/01-introduction.md"


class DrawingTests(BaseTest):
    def test_new_drawing_is_an_empty_excalidraw_scene(self):
        slug = self.create()
        r = self.c.post(self.api(slug, "file/"), {"path": "notes/map.excalidraw"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        item = next(i for i in self.c.get(self.api(slug, "tree/")).data if i["path"] == "notes/map.excalidraw")
        self.assertEqual((item["kind"], item["text"]), ("drawing", True))
        scene = json.loads(self.c.get(self.api(slug, "file/?path=notes/map.excalidraw")).data["content"])
        self.assertEqual((scene["type"], scene["elements"]), ("excalidraw", []))


class SheetTests(BaseTest):
    def setUp(self):
        super().setUp()
        self.slug = self.create()

    def test_csv_roundtrip_keeps_delimiter(self):
        self.c.post(self.api(self.slug, "file/"), {"path": "data/a.csv", "content": "name;age\nAda;36\n"}, format="json")
        data = self.c.get(self.api(self.slug, "sheet/?path=data/a.csv")).data
        self.assertEqual(data["delimiter"], ";")
        self.assertEqual(data["sheets"][0]["cells"], [["name", "age"], ["Ada", "36"]])
        cells = [["name", "age", "note"], ["Ada", "36", "says \"hi\", twice"], ["Bob"]]
        r = self.c.put(self.api(self.slug, "sheet/?path=data/a.csv"),
                       {"sheets": [{"name": "a", "cells": cells}], "delimiter": ";"}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        text = self.c.get(self.api(self.slug, "file/?path=data/a.csv")).data["content"]
        self.assertEqual(text, 'name;age;note\nAda;36;"says ""hi"", twice"\nBob;;\n')

    def test_new_excel_file_and_edits(self):
        r = self.c.post(self.api(self.slug, "file/"), {"path": "data/budget.xlsx"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        data = self.c.get(self.api(self.slug, "sheet/?path=data/budget.xlsx")).data
        self.assertEqual([s["name"] for s in data["sheets"]], ["Sheet1"])
        cells = [["item", "cost"], ["paper", "12.5"], ["ink", "3"], ["total", "=SUM(B2:B3)"],
                 ["date", "2024-03-01"], ["code", "'007"]]
        r = self.c.put(self.api(self.slug, "sheet/?path=data/budget.xlsx"), {"sheets": [
            {"name": "Costs", "orig": "Sheet1", "cells": cells},
            {"name": "Notes", "cells": [["=not a formula"]]},
        ]}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["sheets"][1]["cells"], [["=not a formula"]])
        wb = load_workbook(self.projects / self.slug / "data/budget.xlsx")
        self.assertEqual(wb.sheetnames, ["Costs", "Notes"])
        ws = wb["Costs"]
        self.assertEqual((ws["B2"].value, ws["B3"].value, ws["B4"].value), (12.5, 3, "=SUM(B2:B3)"))
        self.assertEqual(ws["B5"].value.year, 2024)
        self.assertEqual(ws["B6"].value, "007")
        self.assertEqual(r.data["sheets"][0]["cells"], cells)

    def test_excel_edit_keeps_formatting_and_cached_values(self):
        path = self.projects / self.slug / "data" / "styled.xlsx"
        path.parent.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = "Title"
        ws["A1"].font = Font(bold=True)
        ws["A2"], ws["A3"] = 1, 2
        wb.save(path)
        data = self.c.get(self.api(self.slug, "sheet/?path=data/styled.xlsx")).data
        cells = data["sheets"][0]["cells"]
        cells[2][0] = "5"
        self.c.put(self.api(self.slug, "sheet/?path=data/styled.xlsx"),
                   {"sheets": [{"name": "Data", "orig": "Data", "cells": cells}]}, format="json")
        ws = load_workbook(path)["Data"]
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws["A3"].value, 5)

    def test_cell_text_roundtrip(self):
        for text in ["", "abc", "12", "1.5", "TRUE", "'12", "'=x", "2024-01-02", "2024-01-02 10:30:00"]:
            self.assertEqual(sheets.cell_text(sheets.parse_input(text)), text)

    def test_bad_sheet_names_rejected(self):
        self.c.post(self.api(self.slug, "file/"), {"path": "data/b.xlsx"}, format="json")
        r = self.c.put(self.api(self.slug, "sheet/?path=data/b.xlsx"),
                       {"sheets": [{"name": "a/b", "cells": []}]}, format="json")
        self.assertEqual(r.status_code, 400)


class LinkTests(BaseTest):
    def setUp(self):
        super().setUp()
        self.slug = self.create()
        put = lambda path, content: self.c.post(self.api(self.slug, "file/"), {"path": path, "content": content}, format="json")
        put("notes/a.md", "See [the intro](../manuscript/01-introduction.md) and [[ideas]] and [web](https://x.org).\n")
        put("notes/b.tex", "\\input{../notes/a.md}\n\\includegraphics{../attachments/missing.png}\n")
        put("notes/c.html", '<p><a href="a.md#part">a</a></p>\n')
        drawing = {"type": "excalidraw", "elements": [{"id": "1", "link": "../notes/c.html"},
                                                       {"id": "2", "link": "a.md", "isDeleted": True}]}
        put("notes/map.excalidraw", json.dumps(drawing))
        self.c.post(self.api(self.slug, "file/"), {"path": "data/t.xlsx"}, format="json")

    def edges(self):
        g = self.c.get(self.api(self.slug, "graph/")).data
        return {(e["source"], e["target"]): e["kinds"] for e in g["edges"]}, g

    def test_inline_links_found_in_every_format(self):
        edges, g = self.edges()
        self.assertEqual(edges, {
            ("notes/a.md", INTRO): ["inline"],
            ("notes/a.md", "notes/ideas.md"): ["inline"],
            ("notes/b.tex", "notes/a.md"): ["inline"],
            ("notes/c.html", "notes/a.md"): ["inline"],
            ("notes/map.excalidraw", "notes/c.html"): ["inline"],
        })
        paths = {n["path"] for n in g["nodes"]}
        self.assertIn("data/t.xlsx", paths)
        self.assertNotIn("project.json", paths)

    def test_manual_links_and_backlinks_follow_renames(self):
        r = self.c.post(self.api(self.slug, "links/"), {"from": "data/t.xlsx", "to": "notes/a.md"}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["outgoing"], [{"path": "notes/a.md", "title": "A", "kinds": ["manual"]}])
        back = self.c.get(self.api(self.slug, "links/?path=notes/a.md")).data["incoming"]
        self.assertEqual({b["path"] for b in back}, {"data/t.xlsx", "notes/b.tex", "notes/c.html"})
        self.c.post(self.api(self.slug, "move/"), {"from": "notes/a.md", "to": "notes/renamed.md"}, format="json")
        edges, _ = self.edges()
        self.assertEqual(edges[("data/t.xlsx", "notes/renamed.md")], ["manual"])
        self.c.delete(self.api(self.slug, "file/?path=notes/renamed.md"))
        manifest = self.c.get(self.api(self.slug, "manifest/")).data
        self.assertNotIn("links", manifest["files"].get("data/t.xlsx", {}))

    def test_remove_manual_link(self):
        self.c.post(self.api(self.slug, "links/"), {"from": INTRO, "to": "notes/b.tex"}, format="json")
        r = self.c.delete(self.api(self.slug, f"links/?from={INTRO}&to=notes/b.tex"))
        self.assertEqual(r.data["outgoing"], [])

    def test_link_to_missing_file_refused(self):
        r = self.c.post(self.api(self.slug, "links/"), {"from": INTRO, "to": "nope.md"}, format="json")
        self.assertEqual(r.status_code, 404)


class CommentTests(BaseTest):
    def setUp(self):
        super().setUp()
        self.slug = self.create()

    def test_threads_replies_resolve_delete(self):
        url = self.api(self.slug, "comments/")
        anchor = {"type": "text", "quote": "Introduction", "prefix": "# ", "suffix": "\n"}
        r = self.c.post(url, {"path": INTRO, "anchor": anchor, "text": "Better title?"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        thread = r.data
        self.assertEqual(thread["comments"][0]["author"], "Ada")
        tid = thread["id"]
        r = self.c.post(url + f"{tid}/", {"path": INTRO, "text": "Agreed"}, format="json")
        self.assertEqual(len(r.data["comments"]), 2)
        self.c.patch(url + f"{tid}/", {"path": INTRO, "resolved": True}, format="json")
        self.assertEqual(self.c.get(url).data, {})
        threads = self.c.get(url, {"path": INTRO}).data
        self.assertTrue(threads[0]["resolved"])
        # Stored in the repository, hidden from the file tree, visible to git.
        status = self.c.get(self.api(self.slug, "git/status/")).data
        self.assertIn(".dm/comments/manuscript/01-introduction.md.json", [f["path"] for f in status["files"]])
        self.assertFalse(any(i["path"].startswith(".dm") for i in self.c.get(self.api(self.slug, "tree/")).data))
        reply_id = threads[0]["comments"][1]["id"]
        self.c.delete(url + f"{tid}/?path={INTRO}&comment_id={reply_id}")
        self.assertEqual(len(self.c.get(url, {"path": INTRO}).data[0]["comments"]), 1)
        self.c.delete(url + f"{tid}/?path={INTRO}")
        self.assertEqual(self.c.get(url, {"path": INTRO}).data, [])

    def test_comments_follow_rename_and_delete(self):
        url = self.api(self.slug, "comments/")
        self.c.post(url, {"path": "notes/ideas.md", "anchor": {"type": "text", "quote": "Ideas"}, "text": "x"}, format="json")
        self.c.post(self.api(self.slug, "move/"), {"from": "notes", "to": "research"}, format="json")
        self.assertEqual(len(self.c.get(url, {"path": "research/ideas.md"}).data), 1)
        self.assertEqual(self.c.get(url).data, {"research/ideas.md": 1})
        self.c.delete(self.api(self.slug, "file/?path=research/ideas.md"))
        self.assertEqual(self.c.get(url).data, {})

    def test_cell_anchor_and_validation(self):
        url = self.api(self.slug, "comments/")
        r = self.c.post(url, {"path": INTRO, "anchor": {"type": "cell", "sheet": "S", "cell": "B3"}, "text": "?"}, format="json")
        self.assertEqual(r.data["anchor"], {"type": "cell", "sheet": "S", "cell": "B3"})
        r = self.c.post(url, {"path": INTRO, "anchor": {"type": "text", "quote": " "}, "text": "?"}, format="json")
        self.assertEqual(r.status_code, 400)
        r = self.c.post(url, {"path": "../x", "anchor": {"type": "text", "quote": "a"}, "text": "?"}, format="json")
        self.assertEqual(r.status_code, 400)


class CommitGraphTests(BaseTest):
    def test_graph_lists_all_branches_with_labels(self):
        slug = self.create()
        self.c.post(self.api(slug, "git/branches/"), {"name": "alt"}, format="json")
        self.c.put(self.api(slug, "file/?path=notes/ideas.md"), {"content": "alt\n"}, format="json")
        self.c.post(self.api(slug, "git/commit/"), {"message": "On alt"}, format="json")
        self.c.post(self.api(slug, "git/checkout/"), {"name": "main"}, format="json")
        data = self.c.get(self.api(slug, "git/graph/")).data
        self.assertEqual(len(data["commits"]), 2)
        top, first = data["commits"]
        self.assertEqual(top["subject"], "On alt")
        self.assertEqual(top["refs"], [{"name": "alt", "type": "branch"}])
        self.assertIn({"name": "HEAD", "type": "head"}, first["refs"])
        self.assertIn({"name": "main", "type": "branch"}, first["refs"])
        self.assertEqual(data["head"], first["sha"])
